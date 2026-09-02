from docx import Document
from pathlib import Path
import re

SRC = '/Users/deep/Downloads/Jimmy Final Draft.docx'
OUT = '/Users/deep/PhysicsIA/Jimmy_Final_Draft_latex.tex'
TABLE_OUT = '/Users/deep/PhysicsIA/Jimmy_Final_Draft_tables.tex'
IMGDIR = 'Images/JimmyDraft'

doc = Document(SRC)

def tex(s):
    if s is None:
        return ''
    s = str(s)
    repl = {
        '\u03bb': '@@LAMBDA@@', '\u0394': '@@DELTA@@', '\u03b8': '@@THETA@@',
        '\u00b1': '@@PM@@', '\u00d7': r'\\times ', '\u2212': '-', '\u2013': '--', '\u2014': '---',
        '\u2018': "`", '\u2019': "'", '\u201c': '``', '\u201d': "''", '\u2026': r'\\ldots',
        '\u2265': r'\\geq ', '\u2264': r'\\leq ',
    }
    for a,b in repl.items(): s=s.replace(a,b)
    s = s.replace('\\', r'\\textbackslash{}')
    for a,b in [('&',r'\\&'),('%',r'\\%'),('$',r'\\$'),('#',r'\\#'),('_',r'\\_'),
                ('{',r'\\{'),('}',r'\\}'),('~',r'\\textasciitilde{}'),('^',r'\\textasciicircum{}')]:
        s=s.replace(a,b)
    s = s.replace('@@LAMBDA@@', r'\ensuremath{\lambda}').replace('@@DELTA@@', r'\ensuremath{\Delta}')
    s = s.replace('@@THETA@@', r'\ensuremath{\theta}').replace('@@PM@@', r'\ensuremath{\pm}')
    s = s.replace(r'\\textbackslash{}', r'\\textbackslash{}')
    return s.strip()

def rel_image(paragraph):
    out=[]
    for blip in paragraph._p.xpath('.//*[local-name()="blip"]'):
        rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if rid:
            part=doc.part.related_parts.get(rid)
            if part:
                out.append(Path(part.partname).name)
    return out

def cell_latex(cell):
    bits=[]
    for p in cell.paragraphs:
        if p.text.strip(): bits.append(tex(p.text))
        for fn in rel_image(p):
            bits.append(r'\includegraphics[width=\linewidth]{%s/%s}'%(IMGDIR,fn))
    return r' \newline '.join(bits)

def table_latex(table, caption, label):
    n=len(table.columns)
    if n==2: widths=['0.46\\linewidth','0.46\\linewidth']
    elif n==3: widths=['0.29\\linewidth','0.34\\linewidth','0.29\\linewidth']
    elif n==4: widths=['0.13\\linewidth']*4
    elif n==6: widths=['0.13\\linewidth']*6
    elif n==7: widths=['0.28\\linewidth']+['0.105\\linewidth']*6
    elif n==8: widths=['0.11\\linewidth']*8
    else: widths=['0.95\\linewidth']
    spec=''.join('p{%s}'%w for w in widths)
    lines=[r'\begin{longtable}{%s}'%spec,
           r'\caption{%s}\label{%s}\\'% (tex(caption),label),
           r'\hline']
    for ri,row in enumerate(table.rows):
        vals=[cell_latex(c) for c in row.cells]
        lines.append(' & '.join(vals)+r' \\')
        if ri==0: lines.append(r'\hline')
    lines += [r'\hline',r'\end{longtable}', '']
    return '\n'.join(lines)

def figure(fn, caption, label, width='0.9\\linewidth'):
    return '\n'.join([r'\begin{figure}[htbp]',r'\centering',r'\includegraphics[width=%s]{%s/%s}'%(width,IMGDIR,fn),r'\caption{%s}\label{%s}'%(tex(caption),label),r'\end{figure}', ''])

p = doc.paragraphs
parts=[]
parts += [r'% Paste-ready fragment converted from ``Jimmy Final Draft.docx``.',
          r'% Preamble addition required by this fragment: \usepackage{longtable} (graphicx is already used by main.tex).', '',
          r'\section{Research question and variables}',
          tex(p[0].text), '', tex(p[1].text), '', r'\textbf{Research question.} '+tex(p[2].text), '']
parts += [r'\section{Background}', tex(p[3].text), '', tex(p[4].text), '',
          r'\begin{equation} s = \frac{\lambda D}{d} \end{equation}',
          tex(p[6].text), '', tex(p[7].text), '', tex(p[8].text), '',
          tex(p[9].text), '', r'\begin{equation} d\sin\theta=\Delta x \end{equation}',
          tex(p[11].text), '', r'\begin{equation} n\lambda \end{equation}',
          tex(p[13].text), '']
parts += [r'\section{Experimental design}', table_latex(doc.tables[0], 'Variables, methods, and justifications', 'tab:variables')]
parts += [r'\section{Data analysis}', tex(p[18].text), '', table_latex(doc.tables[1], 'Qualitative observations and interpretations', 'tab:qualitative')]
parts += [r'\section{Raw and processed measurements}',
          table_latex(doc.tables[2], 'Measurements for data sets 1-1, 1-2, and 1-3', 'tab:data1'),
          table_latex(doc.tables[3], 'Measurements for data sets 2-1, 2-2, 2-3, and 3-1', 'tab:data2'),
          table_latex(doc.tables[4], 'Measurements for data sets 3-2, 3-3, and 4-3', 'tab:data3'),
          table_latex(doc.tables[5], 'Short-range measurements at 100 cm, 80 cm, and 60 cm', 'tab:data4'),
          table_latex(doc.tables[6], 'Short-range measurements at 40 cm and 20 cm', 'tab:data5')]
parts += [r'\section{Processing requirements and calculations}',
          r'\begin{itemize}',
          r'\item '+tex(p[32].text), r'\item '+tex(p[33].text),
          r'\item '+tex(p[35].text), r'\item '+tex(p[37].text),
          r'\item '+tex(p[39].text), r'\end{itemize}', '',
          r'\begin{equation} \text{Average fringe spacing}=\frac{\sum\text{(all fringe-spacing data points)}}{\text{number of data points}} \end{equation}',
          r'\begin{equation} \pm\text{Uncertainty}=\frac{\text{maximum value}-\text{minimum value}}{2} \end{equation}',
          figure('image6.png','Processed-data graph with error bars and linear trend line.','fig:regression'),
          figure('image7.png','Linear-regression output (equation and fit statistics shown in the source image).','fig:regression-output')]
parts += [r'\section{Evaluation and conclusion}', tex(p[45].text), '', tex(p[46].text), '', tex(p[47].text), '']
parts += [r'\appendix', r'\section{Limitations and proposed improvements}',
          table_latex(doc.tables[7], 'Specific data issues, methodological weaknesses, and realistic improvements', 'tab:limitations'),
          table_latex(doc.tables[8], 'Additional sources of uncertainty', 'tab:uncertainties'),
          figure('image1.jpeg','Original projection on graph paper (image 1).','fig:photo1'),
          figure('image2.jpeg','Original projection on graph paper (image 2).','fig:photo2'),
          figure('image3.png','Software measurement view (image 3).','fig:photo3'),
          figure('image4.png','Interference pattern with marked region (image 4).','fig:photo4'),
          figure('image5.png','Interference pattern showing a shadow (image 5).','fig:photo5'),
          figure('image8.jpg','Long-range projection on graph paper (image 8).','fig:photo8')]

Path(OUT).write_text('\n'.join(parts)+'\n', encoding='utf-8')
table_parts = [
    r'% Table-only fragment converted from ``Jimmy Final Draft.docx``.',
    r'% Preamble addition: \usepackage{longtable} (graphicx is already used by main.tex).',
    '',
    r'\section{Experimental design}', table_latex(doc.tables[0], 'Variables, methods, and justifications', 'tab:variables'),
    r'\section{Data analysis}', table_latex(doc.tables[1], 'Qualitative observations and interpretations', 'tab:qualitative'),
    r'\section{Raw and processed measurements}',
    table_latex(doc.tables[2], 'Measurements for data sets 1-1, 1-2, and 1-3', 'tab:data1'),
    table_latex(doc.tables[3], 'Measurements for data sets 2-1, 2-2, 2-3, and 3-1', 'tab:data2'),
    table_latex(doc.tables[4], 'Measurements for data sets 3-2, 3-3, and 4-3', 'tab:data3'),
    table_latex(doc.tables[5], 'Short-range measurements at 100 cm, 80 cm, and 60 cm', 'tab:data4'),
    table_latex(doc.tables[6], 'Short-range measurements at 40 cm and 20 cm', 'tab:data5'),
    r'\section{Limitations and proposed improvements}',
    table_latex(doc.tables[7], 'Specific data issues, methodological weaknesses, and realistic improvements', 'tab:limitations'),
    table_latex(doc.tables[8], 'Additional sources of uncertainty', 'tab:uncertainties'),
]
Path(TABLE_OUT).write_text('\n'.join(table_parts)+'\n', encoding='utf-8')
print(OUT)
print(TABLE_OUT)
